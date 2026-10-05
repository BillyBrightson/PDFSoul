"""Anything → PDF: Office documents, HTML, Markdown, text, eBooks, XPS and images.

Office files go through LibreOffice when it is installed, which keeps the exact layout.
Without it, .docx, .xlsx, .csv and .pptx still convert with a built-in renderer that keeps the
content (headings, lists, tables, images) but not the precise layout.
"""

from __future__ import annotations

import csv
import html
import io
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from threading import Event

import pymupdf as fitz

from pdfsoul.core import engines
from pdfsoul.core.convert import IMAGE_EXTENSIONS, ImagesToPdfOptions, images_to_pdf
from pdfsoul.core.document import normalise_toc, save_pdf
from pdfsoul.core.safe_io import unique_path
from pdfsoul.core.types import (
    EngineMissing,
    PdfSoulError,
    Progress,
    ProgressFn,
    Result,
    check_cancel,
    no_progress,
)

WORD_EXTENSIONS = {".doc", ".docx", ".docm", ".dot", ".dotx", ".odt", ".ott", ".rtf", ".wpd"}
SHEET_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".ods", ".csv"}
SLIDE_EXTENSIONS = {".ppt", ".pptx", ".pps", ".ppsx", ".odp"}
OFFICE_EXTENSIONS = WORD_EXTENSIONS | SHEET_EXTENSIONS | SLIDE_EXTENSIONS | {".odg"}
WEB_EXTENSIONS = {".html", ".htm", ".xhtml", ".md", ".markdown", ".txt"}
EBOOK_EXTENSIONS = {".epub", ".mobi", ".fb2", ".xps", ".oxps", ".cbz", ".svg"}
BUILTIN_OFFICE = {".docx", ".xlsx", ".xlsm", ".csv", ".pptx"}
ALL_EXTENSIONS = OFFICE_EXTENSIONS | WEB_EXTENSIONS | EBOOK_EXTENSIONS | IMAGE_EXTENSIONS

PAPER = {"a4": (595.0, 842.0), "letter": (612.0, 792.0)}
MM = 72 / 25.4


@dataclass
class ToPdfOptions:
    merge: bool = True  # several inputs → one PDF; otherwise one PDF each, into folder ``output``
    paper: str = "a4"  # for text, HTML, Markdown, eBooks and the built-in Office renderer
    landscape: bool = False
    margin_mm: float = 18.0
    use_libreoffice: bool = True


@dataclass
class _Converted:
    doc: fitz.Document
    notes: list[str] = field(default_factory=list)


def to_pdf(inputs: list[Path], output: Path, opts: ToPdfOptions | None = None,
           progress: ProgressFn = no_progress, cancel: Event | None = None) -> Result:
    """Convert each input; merge them into ``output`` or write ``output/<name>.pdf`` each."""
    opts = opts or ToPdfOptions()
    if not inputs:
        raise PdfSoulError("Add at least one file.")
    paths = [Path(p) for p in inputs]
    for path in paths:
        if path.suffix.lower() not in ALL_EXTENSIONS:
            raise PdfSoulError(f"{path.name}: PDFSoul can't convert {path.suffix or 'this'} "
                               "files to PDF.")
        if not path.is_file():
            raise PdfSoulError(f"File not found: {path}")
    tick = Progress(progress, cancel, len(paths) * 10)
    notes: list[str] = []
    outputs: list[Path] = []
    merged = fitz.open() if opts.merge or len(paths) == 1 else None
    toc: list[list] = []
    with tempfile.TemporaryDirectory(prefix="pdfsoul-topdf-") as tmp_name:
        tmp = Path(tmp_name)
        for n, path in enumerate(paths):
            check_cancel(cancel)
            converted = _convert_one(path, opts, tmp / str(n), cancel)
            notes.extend(converted.notes)
            doc = converted.doc
            if merged is not None:
                start = merged.page_count
                if len(paths) > 1:
                    toc.append([1, path.stem, start + 1])
                toc.extend([[lvl + (len(paths) > 1), title, page + start]
                            for lvl, title, page, *_ in doc.get_toc()])
                merged.insert_pdf(doc)
            else:
                target = unique_path(Path(output) / f"{path.stem}.pdf")
                save_pdf(doc, target)
                outputs.append(target)
            doc.close()
            tick.step(10)
        if merged is not None:
            if toc:
                merged.set_toc(normalise_toc(toc))
            save_pdf(merged, Path(output))
            pages = merged.page_count
            merged.close()
            outputs = [Path(output)]
            message = (f"Converted {paths[0].name} to a {pages}-page PDF." if len(paths) == 1
                       else f"Combined {len(paths)} files into a {pages}-page PDF.")
        else:
            message = f"Converted {len(paths)} files to PDF."
    tick.done()
    if notes:
        message += " " + " ".join(dict.fromkeys(notes))
    return Result(outputs, message, {"notes": notes})


def _convert_one(path: Path, opts: ToPdfOptions, tmp: Path, cancel: Event | None) -> _Converted:
    tmp.mkdir(parents=True, exist_ok=True)
    ext = path.suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        target = tmp / "images.pdf"
        images_to_pdf([path], target, ImagesToPdfOptions(), cancel=cancel)
        return _Converted(fitz.open(target))
    if ext in EBOOK_EXTENSIONS:
        return _Converted(_ebook(path, opts))
    if ext in WEB_EXTENSIONS:
        return _Converted(_web(path, opts, cancel))
    soffice = engines.find("soffice") if opts.use_libreoffice else None
    if soffice is not None:
        return _Converted(fitz.open(_libreoffice(soffice, path, tmp, cancel)))
    if ext in BUILTIN_OFFICE:
        builder = {".docx": _docx_html, ".xlsx": _xlsx_html, ".xlsm": _xlsx_html,
                   ".csv": _csv_html, ".pptx": _pptx_html}[ext]
        document = builder(path, tmp)
        landscape = opts.landscape or document.landscape
        doc = render_html(document.html, opts.paper, landscape, opts.margin_mm, css=OFFICE_CSS,
                          base=tmp, cancel=cancel, slides=document.slides)
        return _Converted(doc, ["Office files used the built-in layout; install LibreOffice "
                                "for an exact match."])
    raise EngineMissing(f"{path.name}: converting {ext} files needs LibreOffice. Install it "
                        "(libreoffice.org), or set its path in Settings.")


# ---------------------------------------------------------------------------------------------
# Engines


def _libreoffice(soffice: Path, path: Path, tmp: Path, cancel: Event | None) -> Path:
    """Headless LibreOffice with a throwaway profile, so an open LibreOffice window is fine."""
    profile = tmp / "lo-profile"
    out = tmp / "lo-out"
    out.mkdir(exist_ok=True)
    engines.run([str(soffice), f"-env:UserInstallation={profile.resolve().as_uri()}",
                 "--headless", "--norestore", "--nolockcheck", "--convert-to", "pdf",
                 "--outdir", str(out), str(path)], cancel=cancel, timeout=600)
    result = out / f"{path.stem}.pdf"
    if not result.is_file():
        raise PdfSoulError(f"LibreOffice couldn't convert {path.name}.")
    return result


def _page_rect(paper: str, landscape: bool) -> fitz.Rect:
    w, h = PAPER.get(paper, PAPER["a4"])
    return fitz.Rect(0, 0, h, w) if landscape else fitz.Rect(0, 0, w, h)


def _ebook(path: Path, opts: ToPdfOptions) -> fitz.Document:
    try:
        src = fitz.open(path)
    except Exception as exc:
        raise PdfSoulError(f"Couldn't read {path.name}: {exc}") from exc
    if src.is_reflowable:
        rect = _page_rect(opts.paper, opts.landscape)
        src.layout(width=rect.width, height=rect.height, fontsize=11)
    toc = src.get_toc()
    doc = fitz.open("pdf", src.convert_to_pdf())
    src.close()
    if toc:
        doc.set_toc(normalise_toc(toc))
    return doc


BASE_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #1c1c22; }
h1 { font-size: 22pt; margin: 0 0 8pt 0; }
h2 { font-size: 16pt; margin: 14pt 0 6pt 0; }
h3 { font-size: 13pt; margin: 12pt 0 4pt 0; }
h4, h5, h6 { font-size: 11pt; margin: 10pt 0 4pt 0; }
p { margin: 0 0 7pt 0; }
table { border-collapse: collapse; margin: 6pt 0 10pt 0; }
th, td { border: 0.5pt solid #9a9aa8; padding: 3pt 5pt; vertical-align: top; }
th { background-color: #eeeef3; font-weight: bold; }
pre { font-family: monospace; font-size: 9pt; background-color: #f3f3f6; padding: 6pt; }
code { font-family: monospace; font-size: 9.5pt; }
blockquote { margin: 6pt 0 6pt 12pt; color: #55555f; }
"""

OFFICE_CSS = BASE_CSS + """
table { font-size: 9pt; }
.sheet-title, .slide-title { color: #55555f; }
"""

TEXT_CSS = """
body { margin: 0; }
pre { font-family: monospace; font-size: 9.5pt; line-height: 1.35; margin: 0; }
"""


def _web(path: Path, opts: ToPdfOptions, cancel: Event | None) -> fitz.Document:
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    ext = path.suffix.lower()
    css = BASE_CSS
    if ext in {".md", ".markdown"}:
        import markdown

        text = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
    elif ext == ".txt":
        text = f"<pre>{html.escape(text)}</pre>"
        css = TEXT_CSS
    return render_html(text, opts.paper, opts.landscape, opts.margin_mm, css=css,
                       base=path.parent, cancel=cancel)


def render_html(markup: str, paper: str = "a4", landscape: bool = False,
                margin_mm: float = 18.0, *, css: str = BASE_CSS, base: Path | None = None,
                cancel: Event | None = None, slides: list[str] | None = None) -> fitz.Document:
    """Lay out HTML onto pages with MuPDF's Story engine. Relative images load from ``base``.

    ``slides``: render each fragment on its own page(s) instead of one flowing document.
    """
    mediabox = _page_rect(paper, landscape)
    margin = margin_mm * MM
    where = fitz.Rect(mediabox.x0 + margin, mediabox.y0 + margin, mediabox.x1 - margin,
                      mediabox.y1 - margin)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    archive = fitz.Archive(str(base)) if base is not None else None
    pages = 0
    for fragment in slides or [markup]:
        story = fitz.Story(html=fragment, user_css=css, archive=archive)
        more = True
        while more:
            check_cancel(cancel)
            device = writer.begin_page(mediabox)
            more, _ = story.place(where)
            story.draw(device)
            writer.end_page()
            pages += 1
            if pages > 20_000:
                raise PdfSoulError("This document is too long to lay out.")
    writer.close()
    return fitz.open("pdf", buf.getvalue())


# ---------------------------------------------------------------------------------------------
# Built-in Office renderers (used when LibreOffice isn't installed)


@dataclass
class _Html:
    html: str
    landscape: bool = False
    slides: list[str] | None = None


def _esc(text: str) -> str:
    return html.escape(text or "").replace("\n", "<br/>")


class _ImageStore:
    """Images pulled out of a document, written beside the HTML so the Story can load them."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self.count = 0

    def add(self, blob: bytes, ext: str = "png", width_pt: float | None = None) -> str:
        self.count += 1
        name = f"img{self.count}.{ext.lstrip('.').lower() or 'png'}"
        (self.folder / name).write_bytes(blob)
        style = f' style="width:{width_pt:.0f}pt"' if width_pt else ""
        return f'<img src="{name}"{style}/>'


def _docx_html(path: Path, tmp: Path) -> _Html:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise PdfSoulError(f"Couldn't read {path.name}: {exc}") from exc
    images = _ImageStore(tmp)
    parts: list[str] = []
    open_list: str | None = None

    def run_html(run) -> str:
        out = []
        for blip in run._element.iter(qn("a:blip")):
            rel = blip.get(qn("r:embed"))
            part = document.part.related_parts.get(rel) if rel else None
            if part is not None:
                extent = next(run._element.iter(qn("wp:extent")), None)
                width = int(extent.get("cx")) / 12700 if extent is not None else None
                ext = Path(getattr(part, "partname", "x.png")).suffix or ".png"
                out.append(images.add(part.blob, ext, width))
        text = _esc(run.text)
        if text:
            if run.bold:
                text = f"<b>{text}</b>"
            if run.italic:
                text = f"<i>{text}</i>"
            if run.underline:
                text = f"<u>{text}</u>"
            out.append(text)
        return "".join(out)

    def paragraph_html(par: Paragraph) -> tuple[str, str | None]:
        style = (par.style.name if par.style is not None else "") or ""
        body = "".join(run_html(r) for r in par.runs) or "&#160;"
        align = {1: "center", 2: "right", 3: "justify"}.get(
            int(par.alignment) if par.alignment is not None else 0)
        attr = f' style="text-align:{align}"' if align else ""
        numbered = par._p.pPr is not None and par._p.pPr.numPr is not None
        if style == "Title":
            return f"<h1{attr}>{body}</h1>", None
        if style.startswith("Heading"):
            level = "".join(c for c in style if c.isdigit()) or "2"
            return f"<h{min(int(level), 6)}{attr}>{body}</h{min(int(level), 6)}>", None
        if "List Number" in style:
            return f"<li>{body}</li>", "ol"
        if "List" in style or numbered:
            return f"<li>{body}</li>", "ul"
        return f"<p{attr}>{body}</p>", None

    def table_html(table: Table) -> str:
        rows = []
        for r, row in enumerate(table.rows):
            tag = "th" if r == 0 else "td"
            cells = []
            for cell in row.cells:
                inner = "<br/>".join("".join(run_html(run) for run in p.runs)
                                     for p in cell.paragraphs)
                cells.append(f"<{tag}>{inner}</{tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        return "<table>" + "".join(rows) + "</table>"

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            markup, list_kind = paragraph_html(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            markup, list_kind = table_html(Table(child, document)), None
        else:
            continue
        if list_kind != open_list:
            if open_list:
                parts.append(f"</{open_list}>")
            if list_kind:
                parts.append(f"<{list_kind}>")
            open_list = list_kind
        parts.append(markup)
    if open_list:
        parts.append(f"</{open_list}>")
    section = document.sections[0] if document.sections else None
    landscape = bool(section and section.page_width and section.page_height
                     and section.page_width > section.page_height)
    return _Html("\n".join(parts), landscape)


MAX_SHEET_ROWS = 5000


def _rows_table(rows: list[list[str]]) -> str:
    out = []
    for r, row in enumerate(rows):
        tag = "th" if r == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{_esc(c)}</{tag}>" for c in row) + "</tr>")
    return "<table>" + "".join(out) + "</table>"


def _trim(rows: list[list[str]]) -> list[list[str]]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    width = max((max((i + 1 for i, c in enumerate(r) if c.strip()), default=0) for r in rows),
                default=0)
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".") if value % 1 else f"{value:,.0f}"
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _xlsx_html(path: Path, _tmp: Path) -> _Html:
    from openpyxl import load_workbook

    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise PdfSoulError(f"Couldn't read {path.name}: {exc}") from exc
    parts, widest = [], 0
    for sheet in book.worksheets:
        rows = []
        for n, row in enumerate(sheet.iter_rows(values_only=True)):
            if n >= MAX_SHEET_ROWS:
                break
            rows.append([_format_cell(v) for v in row])
        rows = _trim(rows)
        if not rows:
            continue
        widest = max(widest, len(rows[0]))
        parts.append(f'<h2 class="sheet-title">{_esc(sheet.title)}</h2>{_rows_table(rows)}')
    book.close()
    if not parts:
        parts.append("<p>(This workbook is empty.)</p>")
    return _Html("\n".join(parts), landscape=widest > 6)


def _csv_html(path: Path, _tmp: Path) -> _Html:
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = _trim([list(r) for _, r in zip(range(MAX_SHEET_ROWS),
                                          csv.reader(io.StringIO(text), dialect), strict=False)])
    if not rows:
        return _Html("<p>(This file is empty.)</p>")
    return _Html(_rows_table(rows), landscape=len(rows[0]) > 6)


def _pptx_html(path: Path, tmp: Path) -> _Html:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        deck = Presentation(str(path))
    except Exception as exc:
        raise PdfSoulError(f"Couldn't read {path.name}: {exc}") from exc
    images = _ImageStore(tmp)
    slides: list[str] = []

    def shape_html(shape, title_id: int | None) -> str:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return "".join(shape_html(s, None) for s in shape.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            width = shape.width / 12700 if shape.width else None
            return images.add(shape.image.blob, shape.image.ext, min(width or 400, 640))
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            return _rows_table(rows)
        if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
            return ""
        if shape.shape_id == title_id:
            return f'<h1 class="slide-title">{_esc(shape.text_frame.text)}</h1>'
        items = [p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs)]
        if len(items) > 1:
            lis = "".join(f'<li style="margin-left:{p.level * 14}pt">'
                          f"{_esc(''.join(r.text for r in p.runs))}</li>" for p in items)
            return f"<ul>{lis}</ul>"
        return f"<p>{_esc(shape.text_frame.text)}</p>"

    for slide in deck.slides:
        title = slide.shapes.title
        title_id = title.shape_id if title is not None else None
        shapes = sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0))
        body = "".join(shape_html(s, title_id) for s in shapes)
        slides.append(body or "<p>&#160;</p>")
    if not slides:
        slides = ["<p>(This presentation has no slides.)</p>"]
    return _Html("", landscape=True, slides=slides)


def supported_description() -> str:
    return "Word, Excel, PowerPoint, OpenDocument, HTML, Markdown, text, ePub, XPS and images"


SUPPORTED_GROUPS: list[tuple[str, set[str]]] = [
    ("Office documents", OFFICE_EXTENSIONS),
    ("Web and text", WEB_EXTENSIONS),
    ("eBooks and XPS", EBOOK_EXTENSIONS),
    ("Images", IMAGE_EXTENSIONS),
]


def file_filter(extensions: set[str] | None = None) -> str:
    """A Qt/OS file-dialog filter, e.g. ``Supported files (*.docx *.html …)``."""
    exts = sorted(extensions or ALL_EXTENSIONS)
    groups = [f"{name} ({' '.join('*' + e for e in sorted(group & set(exts)))})"
              for name, group in SUPPORTED_GROUPS if group & set(exts)]
    return ";;".join([f"Supported files ({' '.join('*' + e for e in exts)})", *groups])
