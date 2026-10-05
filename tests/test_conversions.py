"""PDF → Office/web/PDF-A, anything → PDF, and OCR."""

from __future__ import annotations

from pathlib import Path

import pikepdf
import pymupdf as fitz
import pytest
from conftest import PASSWORD

from pdfsoul.core import engines, export, ocr, to_pdf
from pdfsoul.core.types import EngineMissing, PdfSoulError

needs_gs = pytest.mark.skipif(engines.find("gs") is None, reason="Ghostscript not installed")
needs_tesseract = pytest.mark.skipif(engines.find("tesseract") is None,
                                     reason="Tesseract not installed")


@pytest.fixture(scope="module")
def report(tmp_path_factory) -> Path:
    """A heading, a paragraph, a ruled 3x3 table and a bullet list."""
    path = tmp_path_factory.mktemp("report") / "report.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Annual Report", fontsize=24, fontname="hebo")
    page.insert_text((72, 110), "Revenue grew strongly this year across regions.", fontsize=11)
    rows = [["Region", "Q1", "Q2"], ["North", "1,200", "1,350.50"], ["South", "980", "(12)"]]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = fitz.Rect(72 + c * 120, 140 + r * 24, 192 + c * 120, 164 + r * 24)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_textbox(fitz.Rect(rect.x0 + 4, rect.y0 + 5, rect.x1, rect.y1), value,
                                fontsize=10)
    page.insert_text((72, 260), "• First point", fontsize=11)
    page.insert_text((72, 276), "• Second point", fontsize=11)
    doc.new_page().insert_text((72, 80), "Second page text", fontsize=11)
    doc.save(path)
    return path


# --- PDF → other formats ----------------------------------------------------------------------


def test_pdf_to_word(report, out):
    import docx

    result = export.pdf_to_word([report], out / "r.docx")
    text = "\n".join(p.text for p in docx.Document(str(out / "r.docx")).paragraphs)
    assert "Annual Report" in text and "Revenue grew" in text
    assert "2 pages" in result.message


def test_pdf_to_word_page_range_and_encrypted(fixtures, out):
    import docx

    export.pdf_to_word([fixtures["encrypted"]], out / "e.docx",
                       export.PdfToWordOptions(pages="2", password=PASSWORD))
    text = "\n".join(p.text for p in docx.Document(str(out / "e.docx")).paragraphs)
    assert "Page 2" in text and "Page 1 " not in text


def test_pdf_to_excel_tables_become_numbers(report, out):
    from openpyxl import load_workbook

    result = export.pdf_to_excel([report], out / "r.xlsx")
    book = load_workbook(out / "r.xlsx")
    assert book.sheetnames == ["Page 1 table 1"]
    rows = [[c.value for c in row] for row in book.worksheets[0].iter_rows()]
    assert rows == [["Region", "Q1", "Q2"], ["North", 1200, 1350.5], ["South", 980, -12]]
    assert result.details["tables"] == 1


def test_pdf_to_excel_without_tables_falls_back_to_text(fixtures, out):
    from openpyxl import load_workbook

    result = export.pdf_to_excel([fixtures["simple"]], out / "s.xlsx",
                                 export.PdfToExcelOptions(pages="1-2"))
    assert "No tables" in result.message
    assert load_workbook(out / "s.xlsx").sheetnames == ["Page 1", "Page 2"]


@pytest.mark.parametrize("value,expected", [
    ("1,234.50", 1234.5), ("(12)", -12), ("7%", 0.07), ("-3", -3), ("42", 42),
    ("abc", "abc"), ("", None), ("12 apples", "12 apples"),
])
def test_excel_cell_values(value, expected):
    assert export._cell_value(value) == expected


def test_pdf_to_powerpoint(report, out):
    from pptx import Presentation

    export.pdf_to_powerpoint([report], out / "r.pptx",
                             export.PdfToPowerPointOptions(dpi=100))
    deck = Presentation(str(out / "r.pptx"))
    assert len(deck.slides) == 2
    assert deck.slide_width == int(595 * export.EMU_PER_PT)
    assert "Annual Report" in deck.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("mode", list(export.HtmlMode))
def test_pdf_to_html(report, out, mode):
    export.pdf_to_html([report], out / "r.html", export.PdfToHtmlOptions(mode=mode))
    html = (out / "r.html").read_text(encoding="utf-8")
    assert html.startswith("<!DOCTYPE html>") and "Annual Report" in html
    assert html.count('class="page') == 2


def test_pdf_to_markdown_structure(report, out):
    export.pdf_to_markdown([report], out / "r.md")
    md = (out / "r.md").read_text(encoding="utf-8")
    assert md.startswith("# Annual Report")
    assert "| Region | Q1 | Q2 |" in md and "|---|---|---|" in md
    assert "- First point\n- Second point" in md
    assert "Second page text" in md


@needs_gs
def test_to_pdfa_has_output_intent_and_metadata(report, out):
    result = export.to_pdfa([report], out / "a.pdf",
                            export.PdfaOptions(level=export.PdfaLevel.A2B))
    with pikepdf.open(out / "a.pdf") as pdf:
        assert "/OutputIntents" in pdf.Root
        assert pdf.open_metadata().get("pdfaid:part") == "2"
    assert "PDF/A-2b" in result.message


@needs_gs
def test_to_pdfa_encrypted_input(fixtures, out):
    export.to_pdfa([fixtures["encrypted"]], out / "a.pdf", export.PdfaOptions(password=PASSWORD))
    with pikepdf.open(out / "a.pdf") as pdf:
        assert not pdf.is_encrypted


def test_to_grayscale(fixtures, out):
    export.to_grayscale([fixtures["scanned"]], out / "g.pdf")
    with fitz.open(out / "g.pdf") as doc:
        pix = doc[0].get_pixmap(dpi=30)
        r, g, b = pix.pixel(pix.width // 3, pix.height // 3)
        assert abs(r - g) <= 2 and abs(g - b) <= 2


def test_to_grayscale_without_ghostscript(fixtures, out, monkeypatch):
    monkeypatch.setattr(engines, "find", lambda name: None)
    export.to_grayscale([fixtures["simple"]], out / "g.pdf")
    assert fitz.open(out / "g.pdf").page_count == 10


def test_pdfa_without_ghostscript_is_clear(report, out, monkeypatch):
    monkeypatch.setattr(engines, "find", lambda name: None)
    with pytest.raises(EngineMissing, match="Ghostscript"):
        export.to_pdfa([report], out / "a.pdf")


# --- Anything → PDF ---------------------------------------------------------------------------


@pytest.fixture
def documents(tmp_path) -> dict[str, Path]:
    import docx
    from openpyxl import Workbook
    from pptx import Presentation

    d = tmp_path / "docs"
    d.mkdir()
    files = {}
    files["md"] = d / "note.md"
    files["md"].write_text("# Title\n\nSome *markdown*.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    files["html"] = d / "page.html"
    files["html"].write_text("<h1>Hello</h1><p>World <b>bold</b></p>")
    files["txt"] = d / "plain.txt"
    files["txt"].write_text("plain text\n    indented\n")
    files["csv"] = d / "data.csv"
    files["csv"].write_text("name;amount\nAda;10\nBob;20\n")
    files["svg"] = d / "shape.svg"
    files["svg"].write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
                            '<rect width="200" height="100" fill="red"/></svg>')
    document = docx.Document()
    document.add_heading("Doc heading", 1)
    document.add_paragraph("Body ").add_run("bold").bold = True
    document.add_paragraph("item one", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 1).text = "head", "value"
    files["docx"] = d / "letter.docx"
    document.save(str(files["docx"]))
    book = Workbook()
    book.active.append(["x", "y"])
    book.active.append([1, 2.5])
    files["xlsx"] = d / "sheet.xlsx"
    book.save(files["xlsx"])
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "point a\npoint b"
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Second slide"
    files["pptx"] = d / "deck.pptx"
    deck.save(str(files["pptx"]))
    return files


BUILTIN = to_pdf.ToPdfOptions(use_libreoffice=False)


@pytest.mark.parametrize("kind,expected", [
    ("md", "Title"), ("html", "Hello"), ("txt", "indented"), ("csv", "Ada"),
    ("docx", "Doc heading"), ("xlsx", "2.5"), ("pptx", "Slide title"),
])
def test_to_pdf_each_kind(documents, out, kind, expected):
    result = to_pdf.to_pdf([documents[kind]], out / "o.pdf", BUILTIN)
    with fitz.open(out / "o.pdf") as doc:
        assert expected in "".join(page.get_text() for page in doc)
    assert result.outputs == [out / "o.pdf"]


def test_docx_builtin_keeps_structure(documents, out):
    to_pdf.to_pdf([documents["docx"]], out / "o.pdf", BUILTIN)
    text = fitz.open(out / "o.pdf")[0].get_text()
    assert "Body bold" in text and "item one" in text and "value" in text


def test_pptx_builtin_one_page_per_slide(documents, out):
    to_pdf.to_pdf([documents["pptx"]], out / "o.pdf", BUILTIN)
    with fitz.open(out / "o.pdf") as doc:
        assert doc.page_count == 2
        assert doc[0].rect.width > doc[0].rect.height  # landscape
        assert "Second slide" in doc[1].get_text()


def test_svg_to_pdf(documents, out):
    to_pdf.to_pdf([documents["svg"]], out / "o.pdf")
    with fitz.open(out / "o.pdf") as doc:
        assert round(doc[0].rect.width) == 200  # MuPDF maps SVG user units to points
        pix = doc[0].get_pixmap(dpi=36)
        assert pix.pixel(pix.width // 2, pix.height // 2) == (255, 0, 0)


def test_to_pdf_merges_with_bookmarks(documents, out):
    files = [documents["md"], documents["html"], documents["csv"]]
    result = to_pdf.to_pdf(files, out / "all.pdf", BUILTIN)
    with fitz.open(out / "all.pdf") as doc:
        assert [title for _, title, _ in doc.get_toc()] == ["note", "page", "data"]
    assert "Combined 3 files" in result.message
    assert "LibreOffice" in result.message  # the csv used the built-in renderer


def test_to_pdf_separate_files(documents, out):
    result = to_pdf.to_pdf([documents["md"], documents["txt"]], out / "folder",
                           to_pdf.ToPdfOptions(merge=False))
    assert sorted(p.name for p in result.outputs) == ["note.pdf", "plain.pdf"]


def test_to_pdf_landscape_letter(documents, out):
    to_pdf.to_pdf([documents["txt"]], out / "o.pdf",
                  to_pdf.ToPdfOptions(paper="letter", landscape=True))
    rect = fitz.open(out / "o.pdf")[0].rect
    assert (round(rect.width), round(rect.height)) == (792, 612)


def test_to_pdf_long_text_paginates(tmp_path, out):
    long = tmp_path / "long.txt"
    long.write_text("\n".join(f"Line {n}" for n in range(400)))
    to_pdf.to_pdf([long], out / "o.pdf")
    with fitz.open(out / "o.pdf") as doc:
        assert doc.page_count > 3 and "Line 399" in doc[-1].get_text()


def test_to_pdf_unsupported_and_legacy_formats(tmp_path, out, monkeypatch):
    odd = tmp_path / "x.xyz"
    odd.write_text("?")
    with pytest.raises(PdfSoulError, match="can't convert"):
        to_pdf.to_pdf([odd], out / "o.pdf")
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    monkeypatch.setattr(engines, "find", lambda name: None)
    with pytest.raises(EngineMissing, match="LibreOffice"):
        to_pdf.to_pdf([legacy], out / "o.pdf")


def test_to_pdf_uses_libreoffice_when_present(documents, out, monkeypatch, tmp_path):
    calls = []

    def fake_soffice(cmd, **_kw):
        calls.append(cmd)
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "from libreoffice")
        doc.save(outdir / f"{Path(cmd[-1]).stem}.pdf")
        return ""

    monkeypatch.setattr(engines, "find", lambda name: Path("/fake/soffice"))
    monkeypatch.setattr(engines, "run", fake_soffice)
    result = to_pdf.to_pdf([documents["docx"]], out / "o.pdf")
    assert "--headless" in calls[0] and "--convert-to" in calls[0]
    assert "from libreoffice" in fitz.open(out / "o.pdf")[0].get_text()
    assert "LibreOffice" not in result.message


# --- OCR ---------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def scan(report, tmp_path_factory) -> Path:
    """The report rendered to an image-only page, with a 90° page rotation."""
    path = tmp_path_factory.mktemp("scan") / "scan.pdf"
    pix = fitz.open(report)[0].get_pixmap(dpi=200)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    page = doc.new_page()
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    page.set_rotation(90)
    doc.save(path)
    return path


@needs_tesseract
def test_ocr_makes_scan_searchable(scan, out):
    result = ocr.ocr([scan], out / "o.pdf", ocr.OcrOptions(dpi=200))
    with fitz.open(out / "o.pdf") as doc:
        assert "Annual Report" in doc[0].get_text()
        assert "Annual Report" in doc[1].get_text()
        assert doc[1].rotation == 90
        assert doc[0].search_for("Revenue")
    assert result.details["ocr_pages"] == 2


@needs_tesseract
def test_ocr_skips_pages_with_text(fixtures, out):
    result = ocr.ocr([fixtures["simple"]], out / "o.pdf", ocr.OcrOptions(pages="1-2"))
    assert result.details == {"ocr_pages": 0, "skipped": 2}


@needs_tesseract
def test_ocr_unknown_language(scan, out):
    with pytest.raises(PdfSoulError, match="language pack"):
        ocr.ocr([scan], out / "o.pdf", ocr.OcrOptions(language="klingon"))


def test_ocr_without_tesseract(scan, out, monkeypatch):
    monkeypatch.setattr(engines, "find", lambda name: None)
    with pytest.raises(EngineMissing, match="Tesseract"):
        ocr.ocr([scan], out / "o.pdf")
    assert ocr.languages() == []
